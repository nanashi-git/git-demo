"""
Создание шаблона реферата/доклада (.docx)
Оформление по ГОСТ-подобным стандартам с красивой типографикой
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
import os


def setup_styles(doc):
    """Настройка стилей документа"""
    # Стиль Normal
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.space_after = Pt(0)
    paragraph_format.first_line_indent = Cm(1.25)

    # Стиль Heading 1
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.first_line_indent = Cm(0)

    # Стиль Heading 2
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.first_line_indent = Cm(1.25)

    return doc


def setup_page(doc):
    """Настройка полей страницы (ГОСТ)"""
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
    return doc


def create_title_page(doc):
    """Создание титульного листа"""
    # Министерство / Ведомство
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run('МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ\nРОССИЙСКОЙ ФЕДЕРАЦИИ')
    run.font.size = Pt(14)
    run.font.name = 'Times New Roman'

    # Название ВУЗа
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Cm(0)
    p2.paragraph_format.space_before = Pt(12)
    run2 = p2.add_run('ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ\nОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ')
    run2.font.size = Pt(12)
    run2.font.name = 'Times New Roman'

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.first_line_indent = Cm(0)
    p3.paragraph_format.space_before = Pt(6)
    run3 = p3.add_run('«НАЗВАНИЕ УНИВЕРСИТЕТА»')
    run3.font.size = Pt(14)
    run3.font.bold = True
    run3.font.name = 'Times New Roman'

    # Кафедра
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.first_line_indent = Cm(0)
    p4.paragraph_format.space_before = Pt(24)
    run4 = p4.add_run('Кафедра «Название кафедры»')
    run4.font.size = Pt(14)
    run4.font.name = 'Times New Roman'

    # Отступ перед типом работы
    for _ in range(3):
        doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)

    # Тип работы
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.first_line_indent = Cm(0)
    run5 = p5.add_run('РЕФЕРАТ')
    run5.font.size = Pt(18)
    run5.font.bold = True
    run5.font.name = 'Times New Roman'

    # Дисциплина
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.first_line_indent = Cm(0)
    p6.paragraph_format.space_before = Pt(12)
    run6 = p6.add_run('по дисциплине «Название дисциплины»')
    run6.font.size = Pt(14)
    run6.font.name = 'Times New Roman'

    # Тема
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.first_line_indent = Cm(0)
    p7.paragraph_format.space_before = Pt(12)
    run7 = p7.add_run('на тему: «Тема реферата»')
    run7.font.size = Pt(14)
    run7.font.bold = True
    run7.font.name = 'Times New Roman'

    # Отступ перед информацией о студенте
    for _ in range(4):
        doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)

    # Информация о студенте (справа)
    p8 = doc.add_paragraph()
    p8.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p8.paragraph_format.first_line_indent = Cm(0)
    run8 = p8.add_run('Выполнил: студент группы XX-XXX\nФамилия И.О.')
    run8.font.size = Pt(14)
    run8.font.name = 'Times New Roman'

    p9 = doc.add_paragraph()
    p9.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p9.paragraph_format.first_line_indent = Cm(0)
    p9.paragraph_format.space_before = Pt(12)
    run9 = p9.add_run('Проверил: должность\nФамилия И.О.')
    run9.font.size = Pt(14)
    run9.font.name = 'Times New Roman'

    # Отступ
    for _ in range(3):
        doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)

    # Город и год
    p10 = doc.add_paragraph()
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p10.paragraph_format.first_line_indent = Cm(0)
    run10 = p10.add_run('Город — 2026')
    run10.font.size = Pt(14)
    run10.font.name = 'Times New Roman'

    # Разрыв страницы
    doc.add_page_break()


def create_table_of_contents(doc):
    """Создание содержания (шаблон)"""
    doc.add_heading('СОДЕРЖАНИЕ', level=1)

    items = [
        ('Введение', '3'),
        ('1. Название первого раздела', '5'),
        ('1.1. Подраздел первого раздела', '5'),
        ('1.2. Подраздел первого раздела', '7'),
        ('2. Название второго раздела', '9'),
        ('2.1. Подраздел второго раздела', '9'),
        ('2.2. Подраздел второго раздела', '11'),
        ('3. Название третьего раздела', '13'),
        ('Заключение', '15'),
        ('Список использованных источников', '16'),
    ]

    for title, page in items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), alignment=WD_ALIGN_PARAGRAPH.RIGHT, leader=1)
        run = p.add_run(title)
        run.font.size = Pt(14)
        run.font.name = 'Times New Roman'
        run2 = p.add_run('\t')
        run3 = p.add_run(page)
        run3.font.size = Pt(14)
        run3.font.name = 'Times New Roman'

    doc.add_page_break()


def create_introduction(doc):
    """Создание введения"""
    doc.add_heading('ВВЕДЕНИЕ', level=1)

    intro_text = (
        'Актуальность темы данной работы обусловлена тем, что... '
        '[Здесь необходимо описать, почему данная тема является важной и актуальной '
        'в настоящее время. Рекомендуемый объём введения: 1-2 страницы.]'
    )
    p = doc.add_paragraph(intro_text)

    p2 = doc.add_paragraph()
    run = p2.add_run('Цель работы')
    run.bold = True
    run2 = p2.add_run(' — [сформулировать цель реферата].')

    p3 = doc.add_paragraph()
    run3 = p3.add_run('Задачи работы:')
    run3.bold = True

    tasks = [
        'изучить теоретические аспекты...',
        'проанализировать...',
        'рассмотреть...',
        'сделать выводы о...',
    ]
    for task in tasks:
        p_task = doc.add_paragraph(f'— {task}')
        p_task.paragraph_format.first_line_indent = Cm(1.25)

    doc.add_page_break()


def create_main_body(doc):
    """Создание основной части"""
    doc.add_heading('1. НАЗВАНИЕ ПЕРВОГО РАЗДЕЛА', level=1)

    p = doc.add_paragraph(
        '[Здесь размещается основной текст первого раздела. '
        'Рекомендуется раскрыть теоретические основы рассматриваемой темы. '
        'Текст должен быть связным, логичным и структурированным.]'
    )

    doc.add_heading('1.1. Подраздел первого раздела', level=2)
    doc.add_paragraph(
        '[Текст подраздела. Каждый подраздел раскрывает отдельный аспект '
        'основного раздела.]'
    )

    doc.add_heading('1.2. Подраздел первого раздела', level=2)
    doc.add_paragraph(
        '[Текст второго подраздела.]'
    )

    doc.add_page_break()

    doc.add_heading('2. НАЗВАНИЕ ВТОРОГО РАЗДЕЛА', level=1)
    doc.add_paragraph(
        '[Второй раздел, как правило, содержит практическую часть '
        'или более глубокий анализ темы.]'
    )

    doc.add_page_break()


def create_conclusion(doc):
    """Создание заключения"""
    doc.add_heading('ЗАКЛЮЧЕНИЕ', level=1)

    doc.add_paragraph(
        '[В заключении подводятся итоги работы, формулируются основные выводы. '
        'Заключение должно соответствовать поставленным во введении цели и задачам. '
        'Рекомендуемый объём: 1-2 страницы.]'
    )

    doc.add_page_break()


def create_references(doc):
    """Создание списка литературы"""
    doc.add_heading('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ', level=1)

    references = [
        'Фамилия, И.О. Название книги / И.О. Фамилия. — Город: Издательство, 2023. — 200 с.',
        'Фамилия, И.О. Название статьи // Название журнала. — 2024. — № 1. — С. 10-15.',
        'Название электронного ресурса [Электронный ресурс]. — URL: https://example.com (дата обращения: 01.01.2026).',
        'Фамилия, И.О. Название учебного пособия: учеб. пособие / И.О. Фамилия. — 2-е изд. — Город: Издательство, 2022. — 350 с.',
        'Федеральный закон от 01.01.2020 № 1-ФЗ «О ...» // Российская газета. — 2020. — № 1.',
    ]

    for i, ref in enumerate(references, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        run = p.add_run(f'{i}. {ref}')
        run.font.size = Pt(14)
        run.font.name = 'Times New Roman'


def main():
    doc = Document()

    # Настройка
    setup_page(doc)
    setup_styles(doc)

    # Создание структуры
    create_title_page(doc)
    create_table_of_contents(doc)
    create_introduction(doc)
    create_main_body(doc)
    create_conclusion(doc)
    create_references(doc)

    output_path = os.path.join(os.path.dirname(__file__), "template_referat.docx")
    doc.save(output_path)
    print(f"Шаблон реферата создан: {output_path}")


if __name__ == "__main__":
    main()
