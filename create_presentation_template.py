"""
Создание стильного шаблона презентации (.pptx)
Современный минималистичный дизайн с градиентами
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from copy import deepcopy
import os

# Цветовая палитра (современный тёмно-синий + акценты)
COLORS = {
    'primary': RGBColor(0x1A, 0x1A, 0x2E),       # Тёмно-синий фон
    'secondary': RGBColor(0x16, 0x21, 0x3E),      # Чуть светлее
    'accent': RGBColor(0x0F, 0x8B, 0x8D),         # Бирюзовый акцент
    'accent_light': RGBColor(0x53, 0xCE, 0xD1),   # Светло-бирюзовый
    'white': RGBColor(0xFF, 0xFF, 0xFF),           # Белый
    'light_gray': RGBColor(0xE8, 0xE8, 0xE8),     # Светло-серый
    'dark_gray': RGBColor(0x6C, 0x75, 0x7D),      # Тёмно-серый
    'gradient_start': RGBColor(0x0F, 0x0C, 0x29), # Градиент начало
    'gradient_end': RGBColor(0x30, 0x2B, 0x63),   # Градиент конец
}


def add_gradient_background(slide, color1, color2):
    """Добавляет градиентный фон на слайд"""
    background = slide.background
    fill = background.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = color1
    fill.gradient_stops[1].color.rgb = color2


def add_solid_background(slide, color):
    """Добавляет сплошной фон"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_line(slide, left, top, width, height, color):
    """Добавляет акцентную линию"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_decorative_circle(slide, left, top, size, color, transparency=0):
    """Добавляет декоративный круг"""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def create_title_slide(prs):
    """Создаёт титульный слайд"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # Градиентный фон
    add_gradient_background(slide, COLORS['gradient_start'], COLORS['gradient_end'])

    # Декоративные элементы
    add_decorative_circle(slide, Inches(8.5), Inches(-1), Inches(3), COLORS['accent'])
    add_decorative_circle(slide, Inches(-0.5), Inches(5.5), Inches(2), COLORS['accent_light'])

    # Акцентная линия сверху
    add_accent_line(slide, Inches(1.5), Inches(2.8), Inches(2), Pt(4), COLORS['accent'])

    # Заголовок
    txBox = slide.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(7), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "НАЗВАНИЕ ПРЕЗЕНТАЦИИ"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']

    # Подзаголовок
    txBox2 = slide.shapes.add_textbox(Inches(1.5), Inches(4.5), Inches(7), Inches(1))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Подзаголовок или краткое описание темы"
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLORS['light_gray']

    # Автор и дата
    txBox3 = slide.shapes.add_textbox(Inches(1.5), Inches(6.2), Inches(7), Inches(0.8))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "Автор: Имя Фамилия  |  Группа: XX-XXX  |  2026"
    p3.font.size = Pt(12)
    p3.font.color.rgb = COLORS['dark_gray']


def create_section_slide(prs):
    """Создаёт слайд-разделитель"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_gradient_background(slide, COLORS['primary'], COLORS['secondary'])

    # Большой номер секции
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(2), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "01"
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = COLORS['accent']

    # Акцентная линия
    add_accent_line(slide, Inches(1), Inches(3.8), Inches(3), Pt(3), COLORS['accent'])

    # Название раздела
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(8), Inches(1.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "НАЗВАНИЕ РАЗДЕЛА"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = COLORS['white']


def create_content_slide(prs):
    """Создаёт слайд с контентом"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_solid_background(slide, COLORS['primary'])

    # Заголовок слайда
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Заголовок слайда"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']

    # Акцентная линия под заголовком
    add_accent_line(slide, Inches(0.8), Inches(1.2), Inches(1.5), Pt(3), COLORS['accent'])

    # Основной текст
    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True

    bullets = [
        "Первый ключевой пункт вашей презентации",
        "Второй важный момент для аудитории",
        "Третий пункт с дополнительной информацией",
        "Четвёртый пункт — выводы или примеры",
    ]

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = f"  {bullet}"
        p.font.size = Pt(16)
        p.font.color.rgb = COLORS['light_gray']
        p.space_before = Pt(16)


def create_two_column_slide(prs):
    """Создаёт слайд с двумя колонками"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_solid_background(slide, COLORS['primary'])

    # Заголовок
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Сравнение / Два аспекта"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']

    add_accent_line(slide, Inches(0.8), Inches(1.2), Inches(1.5), Pt(3), COLORS['accent'])

    # Левая колонка
    left_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.6), Inches(4.3), Inches(5.2))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = COLORS['secondary']
    left_box.line.fill.background()

    txBox_l = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(3.9), Inches(4.8))
    tf_l = txBox_l.text_frame
    tf_l.word_wrap = True
    p_l = tf_l.paragraphs[0]
    p_l.text = "Аспект 1"
    p_l.font.size = Pt(20)
    p_l.font.bold = True
    p_l.font.color.rgb = COLORS['accent_light']
    p_l2 = tf_l.add_paragraph()
    p_l2.text = "\nОписание первого аспекта или колонки. Здесь можно разместить ключевые тезисы."
    p_l2.font.size = Pt(14)
    p_l2.font.color.rgb = COLORS['light_gray']

    # Правая колонка
    right_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.2), Inches(1.6), Inches(4.3), Inches(5.2))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLORS['secondary']
    right_box.line.fill.background()

    txBox_r = slide.shapes.add_textbox(Inches(5.5), Inches(1.9), Inches(3.9), Inches(4.8))
    tf_r = txBox_r.text_frame
    tf_r.word_wrap = True
    p_r = tf_r.paragraphs[0]
    p_r.text = "Аспект 2"
    p_r.font.size = Pt(20)
    p_r.font.bold = True
    p_r.font.color.rgb = COLORS['accent_light']
    p_r2 = tf_r.add_paragraph()
    p_r2.text = "\nОписание второго аспекта или колонки. Сравнительная информация или альтернатива."
    p_r2.font.size = Pt(14)
    p_r2.font.color.rgb = COLORS['light_gray']


def create_quote_slide(prs):
    """Создаёт слайд с цитатой"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_gradient_background(slide, COLORS['primary'], COLORS['gradient_end'])

    # Кавычка (большая декоративная)
    txBox_q = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(1.5), Inches(1.5))
    tf_q = txBox_q.text_frame
    p_q = tf_q.paragraphs[0]
    p_q.text = "\u201C"
    p_q.font.size = Pt(120)
    p_q.font.color.rgb = COLORS['accent']

    # Текст цитаты
    txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.8), Inches(7), Inches(2.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Здесь можно разместить важную цитату, ключевую мысль или определение, которое подчёркивает тему вашей презентации."
    p.font.size = Pt(22)
    p.font.italic = True
    p.font.color.rgb = COLORS['white']

    # Автор цитаты
    txBox2 = slide.shapes.add_textbox(Inches(1.5), Inches(5.5), Inches(7), Inches(0.5))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "\u2014 Автор цитаты"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLORS['accent_light']


def create_conclusion_slide(prs):
    """Создаёт финальный слайд"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_gradient_background(slide, COLORS['gradient_start'], COLORS['gradient_end'])

    # Декоративный круг
    add_decorative_circle(slide, Inches(7.5), Inches(4.5), Inches(3.5), COLORS['accent'])

    # Текст "Спасибо"
    txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(7), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "СПАСИБО ЗА ВНИМАНИЕ"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']

    # Контакты
    txBox2 = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(6), Inches(1.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Вопросы?"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLORS['accent_light']
    p3 = tf2.add_paragraph()
    p3.text = "\nemail@example.com"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLORS['light_gray']


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Создаём все типы слайдов
    create_title_slide(prs)
    create_section_slide(prs)
    create_content_slide(prs)
    create_two_column_slide(prs)
    create_quote_slide(prs)
    create_conclusion_slide(prs)

    output_path = os.path.join(os.path.dirname(__file__), "template_presentation.pptx")
    prs.save(output_path)
    print(f"Шаблон презентации создан: {output_path}")


if __name__ == "__main__":
    main()
