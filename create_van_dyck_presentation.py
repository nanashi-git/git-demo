#!/usr/bin/env python3
"""
Создание презентации: Антонис Ван Дейк — избранные картины
Предмет: История искусств
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from datetime import datetime
import os

# Тема: Искусство / Гуманитарные — тёплая палитра
COLORS = {
    "bg_primary": RGBColor(0x1A, 0x0A, 0x2E),
    "bg_secondary": RGBColor(0x2D, 0x13, 0x4A),
    "accent": RGBColor(0xC9, 0x9A, 0x2E),        # Золотой
    "accent_light": RGBColor(0xE8, 0xC5, 0x6D),   # Светло-золотой
    "text_primary": RGBColor(0xFF, 0xFF, 0xFF),
    "text_secondary": RGBColor(0xE8, 0xE0, 0xD0),
    "text_muted": RGBColor(0x9A, 0x8A, 0x7A),
    "gradient_start": RGBColor(0x0F, 0x07, 0x1A),
    "gradient_end": RGBColor(0x2D, 0x1B, 0x4E),
    "card_bg": RGBColor(0x22, 0x10, 0x3A),
}

FONTS = {
    "heading": "Georgia",
    "body": "Palatino Linotype",
    "fallback": "Times New Roman",
}


def add_gradient_bg(slide, c1, c2):
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = c1
    fill.gradient_stops[1].color.rgb = c2


def add_solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, color):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, left, top, width, height, text, font_name=None, size=Pt(16),
             color=None, bold=False, italic=False, align=PP_ALIGN.LEFT):
    if font_name is None:
        font_name = FONTS["body"]
    if color is None:
        color = COLORS["text_secondary"]
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = size
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.alignment = align
    return tf


def add_multiline_text(slide, left, top, width, height, lines, font_name=None,
                       size=Pt(14), color=None, spacing=Pt(8)):
    """Добавляет многострочный текст с буллетами"""
    if font_name is None:
        font_name = FONTS["body"]
    if color is None:
        color = COLORS["text_secondary"]
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.name = font_name
        p.font.size = size
        p.font.color.rgb = color
        p.space_before = spacing
    return tf


def create_slide_number(slide, num):
    add_text(slide, Inches(9.2), Inches(7.0), Inches(0.6), Inches(0.4),
             str(num), size=Pt(9), color=COLORS["text_muted"], align=PP_ALIGN.RIGHT)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ==========================================
    # СЛАЙД 1: ТИТУЛЬНЫЙ
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_bg(slide, COLORS["gradient_start"], COLORS["gradient_end"])

    # Декор
    add_shape(slide, MSO_SHAPE.OVAL, Inches(8.0), Inches(-1.5), Inches(4), Inches(4), COLORS["accent"])
    add_shape(slide, MSO_SHAPE.OVAL, Inches(-1.2), Inches(5.5), Inches(2.5), Inches(2.5), COLORS["accent_light"])

    # Линия-акцент
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.4), Inches(3), Pt(4), COLORS["accent"])

    # Вуз
    add_text(slide, Inches(1.2), Inches(0.4), Inches(8), Inches(0.5),
             "Тверской филиал РГУ им. А.Н. Косыгина", size=Pt(10), color=COLORS["text_muted"])

    # Дисциплина
    add_text(slide, Inches(1.2), Inches(1.8), Inches(7), Inches(0.5),
             "История искусств", size=Pt(14), color=COLORS["accent_light"], italic=True)

    # Заголовок
    add_text(slide, Inches(1.2), Inches(2.7), Inches(7.5), Inches(1.8),
             "АНТОНИС ВАН ДЕЙК", font_name=FONTS["heading"],
             size=Pt(38), color=COLORS["text_primary"], bold=True)

    # Подзаголовок
    add_text(slide, Inches(1.2), Inches(4.3), Inches(7.5), Inches(0.8),
             "Избранные произведения великого фламандского портретиста",
             size=Pt(16), color=COLORS["text_secondary"])

    # Автор
    add_text(slide, Inches(1.2), Inches(6.3), Inches(7.5), Inches(0.5),
             f"Студент 1 курса  •  {datetime.now().year}",
             size=Pt(11), color=COLORS["text_muted"])

    # ==========================================
    # СЛАЙД 2: О ХУДОЖНИКЕ
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Антонис Ван Дейк (1599–1641)", font_name=FONTS["heading"],
             size=Pt(26), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Фламандский живописец эпохи барокко, ученик Питера Пауля Рубенса",
        "• Родился 22 марта 1599 г. в Антверпене, умер 9 декабря 1641 г. в Лондоне",
        "• Придворный художник английского короля Карла I с 1632 года",
        "• Крупнейший мастер парадного аристократического портрета XVII века",
        "• Работал в Антверпене, Генуе, Риме, Палермо и Лондоне",
        "• Оказал огромное влияние на английскую портретную живопись на 150 лет вперёд",
        "• Создал более 900 картин за свою короткую жизнь (42 года)",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.5),
                       lines, size=Pt(15), spacing=Pt(14))
    create_slide_number(slide, 2)

    # ==========================================
    # СЛАЙД 3: Семейный портрет
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Семейный портрет (ок. 1620–1621)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. 113,5 × 93,5 см",
        "• Хранится в Государственном Эрмитаже, Санкт-Петербург",
        "• Изображена супружеская пара с дочерью на коленях у матери",
        "• Компактная композиция создаёт атмосферу доверительной",
        "  близости и согласия между членами семьи",
        "• Написана в ранний период творчества — Ван Дейку было ~21 год",
        "• Демонстрирует влияние Рубенса в колористике и свободе мазка",
        "• Одна из самых проникновенных работ молодого мастера",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(14), spacing=Pt(12))
    create_slide_number(slide, 3)

    # ==========================================
    # СЛАЙД 4: Св. Мартин и нищий
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Святой Мартин и нищий (ок. 1618)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Дерево, масло. 172 × 158 см",
        "• Алтарный образ церкви Св. Мартина в Завентеме (Бельгия)",
        "• Сюжет: римский воин Мартин Турский разрезает свой плащ,",
        "  чтобы поделиться с замерзающим нищим холодным зимним утром",
        "• Работа написана, когда Ван Дейку было всего ~19 лет",
        "• Заказ изначально предназначался для Рубенса, но был",
        "  передан его талантливому ученику",
        "• Драматичная барочная композиция с сильной диагональю",
        "• Свидетельство раннего мастерства в монументальной живописи",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(14), spacing=Pt(12))
    create_slide_number(slide, 4)

    # ==========================================
    # СЛАЙД 5: Портрет кардинала Бентивольо
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Портрет кардинала Гвидо Бентивольо (ок. 1623)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. Галерея Палатина, Палаццо Питти, Флоренция",
        "• Написан во время пребывания Ван Дейка в Италии (1621–1627)",
        "• Кардинал Бентивольо имел связи с родной Фландрией художника",
        "• «Весь Рим устремился смотреть это чудо искусства, и каждый",
        "  хотел быть написанным рукой нашего художника» (XVIII в.)",
        "• Работа вдохновлена итальянскими мастерами, особенно Тицианом",
        "• Утвердила Ван Дейка как ведущего портретиста своего времени",
        "• Передаёт жизнь, ум и чувственность натуры кардинала",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(14), spacing=Pt(12))
    create_slide_number(slide, 5)

    # ==========================================
    # СЛАЙД 6: Автопортрет 1620
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Автопортрет (ок. 1620–1621)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. 119,7 × 87,9 см",
        "• Хранится в Метрополитен-музее, Нью-Йорк",
        "• Написан предположительно зимой 1620–1621 гг. в Лондоне",
        "• Ван Дейк изобразил себя как светского джентльмена в изысканной",
        "  одежде, без палитры и кистей — атрибутов ремесла",
        "• Небрежная поза руки у подбородка подчёркивает аристократизм",
        "• Отец художника был богатым торговцем тканями — отсюда роскошь",
        "  костюма и самоуверенность 21-летнего мастера",
        "• Виртуозная кисть передаёт блеск ткани и сияние молодой кожи",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(14), spacing=Pt(12))
    create_slide_number(slide, 6)

    # ==========================================
    # СЛАЙД 7: Портрет Марии Луизы де Тассис
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Портрет Марии Луизы де Тассис (ок. 1629–1630)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. Коллекция князей Лихтенштейн, Вадуц",
        "• Мария Луиза (1611–1638) — дочь Антонио де Тассис из Антверпена",
        "• Семья де Тассис из Бергамо создала первую почтовую систему Европы",
        "  в конце XV века (ныне известна как Thurn und Taxis)",
        "• Изображена в возрасте около 19 лет",
        "• Изысканный барочный женский портрет: богатое платье, кружева,",
        "  жемчуг и утончённые черты молодой аристократки",
        "• Яркий пример мастерства Ван Дейка в передаче тканей и украшений",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(14), spacing=Pt(12))
    create_slide_number(slide, 7)

    # ==========================================
    # СЛАЙД 8: Портрет сэра Томаса Чалонера
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Портрет сэра Томаса Чалонера (конец 1630-х)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. Государственный Эрмитаж, Санкт-Петербург",
        "• Одна из последних работ Ван Дейка",
        "• Сэр Томас Чалонер — английский придворный при Карле I",
        "• Художник с мастерством передаёт стареющее лицо с дряблой кожей",
        "  и румянцем на щеках — честная, без лести, характеристика",
        "• Считается одним из лучших полотен позднего периода",
        "• Глубокий психологизм: зритель ощущает характер и жизненный",
        "  опыт изображённого человека",
        "• Свободная, уверенная манера письма зрелого мастера",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(14), spacing=Pt(12))
    create_slide_number(slide, 8)

    # ==========================================
    # СЛАЙД 9: Карл I на охоте
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Портрет Карла I на охоте (ок. 1635)", font_name=FONTS["heading"],
             size=Pt(24), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. 266 × 207 см. Музей Лувр, Париж",
        "• Шедевр Ван Дейка и жемчужина коллекции Лувра с 1793 года",
        "• Карл I изображён в гражданской одежде, стоящим у лошади,",
        "  как будто отдыхающим на охоте",
        "• «Тонкий компромисс между джентльменской небрежностью",
        "  и королевской уверенностью» (описание Лувра)",
        "• Король невысокого роста (163 см), но Ван Дейк с помощью",
        "  композиции создаёт впечатление величественной фигуры",
        "• Лошадь словно кланяется, подчёркивая статус монарха",
        "• Прекрасный пейзажный фон объединяет природу и власть",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(13), spacing=Pt(10))
    create_slide_number(slide, 9)

    # ==========================================
    # СЛАЙД 10: Портрет Дигби и Рассела
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, COLORS["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
             "Портрет Джорджа Дигби и Уильяма Рассела (ок. 1637)", font_name=FONTS["heading"],
             size=Pt(22), color=COLORS["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.05), Inches(2), Pt(3), COLORS["accent"])

    lines = [
        "• Холст, масло. Коллекция Спенсеров, Олторп, Англия",
        "• Двойной парадный портрет двух молодых английских аристократов",
        "• Джордж Дигби (1612–1677) — 2-й граф Бристоль, политик-роялист,",
        "  государственный секретарь Карла I, поэт и драматург",
        "• Уильям Рассел (1616–1700) — впоследствии 1-й герцог Бедфорд",
        "• Оба — молодые придворные, им около 21–25 лет на портрете",
        "• Элегантная непринуждённость поз передаёт светскость",
        "  и уверенность представителей высшей аристократии",
        "• Типичный пример «friendship portrait» — жанра, в котором",
        "  Ван Дейк был непревзойдённым мастером",
    ]
    add_multiline_text(slide, Inches(0.8), Inches(1.4), Inches(8.4), Inches(5.5),
                       lines, size=Pt(13), spacing=Pt(10))
    create_slide_number(slide, 10)

    # ==========================================
    # СЛАЙД 11: ФИНАЛЬНЫЙ
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_bg(slide, COLORS["gradient_start"], COLORS["gradient_end"])

    add_shape(slide, MSO_SHAPE.OVAL, Inches(7.0), Inches(4.0), Inches(4), Inches(4), COLORS["accent"])
    add_shape(slide, MSO_SHAPE.OVAL, Inches(-0.8), Inches(-0.8), Inches(2), Inches(2), COLORS["accent_light"])

    add_text(slide, Inches(1.2), Inches(2.5), Inches(7), Inches(1.5),
             "СПАСИБО ЗА ВНИМАНИЕ", font_name=FONTS["heading"],
             size=Pt(34), color=COLORS["text_primary"], bold=True)

    add_text(slide, Inches(1.2), Inches(4.2), Inches(6), Inches(0.8),
             "Готов ответить на вопросы", size=Pt(18), color=COLORS["accent_light"])

    add_text(slide, Inches(1.2), Inches(5.8), Inches(6), Inches(0.8),
             "Тверской филиал РГУ им. А.Н. Косыгина  •  История искусств",
             size=Pt(11), color=COLORS["text_muted"])

    # Сохранение
    output = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "Van_Dyck_Presentation.pptx")
    prs.save(output)
    print(f"Презентация создана: {output}")
    print(f"  Слайдов: 11")
    print(f"  Тема: Антонис Ван Дейк — избранные произведения")


if __name__ == "__main__":
    main()
