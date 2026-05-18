#!/usr/bin/env python3
"""
ГЕНЕРАТОР ПРЕЗЕНТАЦИЙ И РЕФЕРАТОВ
==================================
Адаптивный инструмент для создания:
- Стильных презентаций (.pptx) с автоматическим подбором дизайна под тему
- Рефератов/докладов (.docx) по ГОСТ 7.32-2017

Для студента 1 курса Тверского филиала РГУ им. А.Н. Косыгина

Использование:
    python generator.py presentation "Тема" --author "Имя Фамилия" --group "XX-XXX"
    python generator.py referat "Тема" --author "Имя Фамилия" --group "XX-XXX"
    python generator.py themes  # Показать доступные темы дизайна
"""

import argparse
import os
import sys
from datetime import datetime

from themes import (
    THEMES, detect_theme, get_theme, list_themes,
    UNIVERSITY_INFO, GOST_STANDARDS,
)


# =============================================================================
# ГЕНЕРАТОР ПРЕЗЕНТАЦИЙ
# =============================================================================

def create_presentation(
    title,
    author="Студент",
    group="XX-XXX",
    department="",
    discipline="",
    slides_count=8,
    subtitle="",
    theme_key=None,
    output=None,
):
    """
    Создаёт стильную презентацию с адаптивным дизайном.
    Тема автоматически подбирается по названию, или можно указать вручную.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    # Определение темы
    if theme_key and theme_key in THEMES:
        theme = get_theme(theme_key)
    else:
        detected = detect_theme(title, subtitle + " " + discipline)
        theme = get_theme(detected)
        theme_key = detected

    colors = theme["colors"]
    fonts = theme["fonts"]
    style = theme["style"]

    print(f"  Тема дизайна: {theme['name']} ({theme_key})")

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # --- Утилиты ---
    def add_gradient_bg(slide, c1, c2):
        fill = slide.background.fill
        fill.gradient()
        fill.gradient_stops[0].color.rgb = c1
        fill.gradient_stops[1].color.rgb = c2

    def add_solid_bg(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_shape(slide, shape_type, left, top, width, height, color):
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    def add_text(slide, left, top, width, height, text, font_name, size,
                 color, bold=False, italic=False, align=PP_ALIGN.LEFT):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = font_name
        p.font.size = size
        p.font.bold = bold
        p.font.italic = italic
        p.font.color.rgb = color
        p.alignment = align
        return tf

    # === СЛАЙД 1: ТИТУЛЬНЫЙ ===
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_bg(slide, colors["gradient_start"], colors["gradient_end"])

    # Декоративные элементы
    add_shape(slide, MSO_SHAPE.OVAL, Inches(8.2), Inches(-1.2), Inches(3.5), Inches(3.5), colors["accent"])
    add_shape(slide, MSO_SHAPE.OVAL, Inches(-1), Inches(5.8), Inches(2.2), Inches(2.2), colors["accent_light"])

    # Акцентная линия
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.6), Inches(2.5), Pt(4), colors["accent"])

    # Вуз (маленький текст сверху)
    add_text(slide, Inches(1.2), Inches(0.4), Inches(7.5), Inches(0.6),
             f"{UNIVERSITY_INFO['branch']} {UNIVERSITY_INFO['short_name']}",
             fonts["body"], Pt(10), colors["text_muted"])

    # Заголовок
    add_text(slide, Inches(1.2), Inches(2.9), Inches(7.5), Inches(1.8),
             title.upper(), fonts["heading"], Pt(34), colors["text_primary"], bold=True)

    # Подзаголовок / дисциплина
    sub = subtitle if subtitle else (f"Дисциплина: {discipline}" if discipline else "")
    if sub:
        add_text(slide, Inches(1.2), Inches(4.7), Inches(7.5), Inches(0.8),
                 sub, fonts["body"], Pt(16), colors["text_secondary"])

    # Автор
    author_line = f"{author}  •  Группа {group}  •  {datetime.now().year}"
    add_text(slide, Inches(1.2), Inches(6.3), Inches(7.5), Inches(0.6),
             author_line, fonts["body"], Pt(11), colors["text_muted"])

    # === СЛАЙД 2: СОДЕРЖАНИЕ ===
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_solid_bg(slide, colors["bg_primary"])

    add_text(slide, Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.9),
             "СОДЕРЖАНИЕ", fonts["heading"], Pt(28), colors["text_primary"], bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(1.8), Pt(3), colors["accent"])

    # Пункты содержания
    toc_items = [
        "01  Введение и актуальность",
        "02  Основная часть",
        "03  Анализ и результаты",
        "04  Выводы",
    ]
    y_pos = 1.8
    for item in toc_items:
        add_text(slide, Inches(0.8), Inches(y_pos), Inches(8), Inches(0.6),
                 item, fonts["body"], Pt(18), colors["text_secondary"])
        y_pos += 0.9

    # === СЛАЙДЫ 3..N-1: КОНТЕНТ ===
    content_slides = slides_count - 3  # минус титульный, содержание и финальный
    for i in range(content_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_solid_bg(slide, colors["bg_primary"])

        # Заголовок
        add_text(slide, Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.9),
                 f"Заголовок слайда {i + 3}",
                 fonts["heading"], Pt(26), colors["text_primary"], bold=True)

        # Акцентная линия
        add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(1.5), Pt(3), colors["accent"])

        # Область для контента
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "[Здесь размещается контент слайда]"
        p.font.name = fonts["body"]
        p.font.size = Pt(16)
        p.font.color.rgb = colors["text_secondary"]

        # Номер слайда внизу
        add_text(slide, Inches(9.0), Inches(7.0), Inches(0.8), Inches(0.4),
                 str(i + 3), fonts["body"], Pt(10), colors["text_muted"],
                 align=PP_ALIGN.RIGHT)

    # === ФИНАЛЬНЫЙ СЛАЙД ===
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_bg(slide, colors["gradient_start"], colors["gradient_end"])

    add_shape(slide, MSO_SHAPE.OVAL, Inches(7.0), Inches(4.0), Inches(4), Inches(4), colors["accent"])

    add_text(slide, Inches(1.2), Inches(2.5), Inches(7), Inches(1.5),
             "СПАСИБО ЗА ВНИМАНИЕ", fonts["heading"], Pt(34),
             colors["text_primary"], bold=True)

    add_text(slide, Inches(1.2), Inches(4.2), Inches(6), Inches(0.8),
             "Готов ответить на вопросы", fonts["body"], Pt(18), colors["accent_light"])

    add_text(slide, Inches(1.2), Inches(5.5), Inches(6), Inches(0.6),
             f"{author} • {UNIVERSITY_INFO['branch']}",
             fonts["body"], Pt(12), colors["text_muted"])

    # === СОХРАНЕНИЕ ===
    if not output:
        safe_title = "".join(c if c.isalnum() or c in ' _-' else '' for c in title)[:50].strip()
        output = f"Презентация - {safe_title}.pptx"

    prs.save(output)
    print(f"\n  Презентация создана: {output}")
    print(f"  Тема: {title}")
    print(f"  Слайдов: {slides_count}")
    print(f"  Автор: {author}")
    print(f"  Дизайн: {theme['name']}")
    return output



# =============================================================================
# ГЕНЕРАТОР РЕФЕРАТОВ
# =============================================================================

def create_referat(
    title,
    author="Студент",
    group="XX-XXX",
    department="",
    discipline="",
    supervisor="",
    supervisor_title="",
    output=None,
):
    """
    Создаёт реферат с оформлением по ГОСТ 7.32-2017.
    Включает данные Тверского филиала РГУ им. А.Н. Косыгина.
    """
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    gost = GOST_STANDARDS
    uni = UNIVERSITY_INFO

    # --- Настройка полей ---
    for section in doc.sections:
        section.top_margin = Cm(gost["margin_top"] / 10)
        section.bottom_margin = Cm(gost["margin_bottom"] / 10)
        section.left_margin = Cm(gost["margin_left"] / 10)
        section.right_margin = Cm(gost["margin_right"] / 10)

    # --- Настройка стилей ---
    style_normal = doc.styles['Normal']
    style_normal.font.name = gost["font_main"]
    style_normal.font.size = Pt(gost["font_size_main"])
    style_normal.paragraph_format.line_spacing = gost["line_spacing"]
    style_normal.paragraph_format.space_after = Pt(0)
    style_normal.paragraph_format.first_line_indent = Cm(gost["first_line_indent"] / 10)

    h1 = doc.styles['Heading 1']
    h1.font.name = gost["font_main"]
    h1.font.size = Pt(gost["font_size_heading1"])
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.first_line_indent = Cm(0)

    h2 = doc.styles['Heading 2']
    h2.font.name = gost["font_main"]
    h2.font.size = Pt(gost["font_size_heading2"])
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.first_line_indent = Cm(gost["first_line_indent"] / 10)

    year = datetime.now().year

    # === ТИТУЛЬНЫЙ ЛИСТ ===
    def add_centered(text, size=14, bold=False, space_before=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        if space_before:
            p.paragraph_format.space_before = Pt(space_before)
        run = p.add_run(text)
        run.font.name = gost["font_main"]
        run.font.size = Pt(size)
        run.font.bold = bold
        return p

    add_centered(uni["ministry"], 12)
    add_centered(uni["branch_full"], 12, space_before=12)

    if department:
        add_centered(f'\nКафедра «{department}»', 14, space_before=18)

    # Отступы
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)

    add_centered('РЕФЕРАТ', 18, bold=True)

    if discipline:
        add_centered(f'по дисциплине «{discipline}»', 14, space_before=12)

    add_centered(f'на тему: «{title}»', 14, bold=True, space_before=12)

    # Отступы перед автором
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)

    # Автор (справа)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(f'Выполнил: студент группы {group}\n{author}')
    run.font.name = gost["font_main"]
    run.font.size = Pt(14)

    # Проверил
    if supervisor:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(12)
        sup_text = f'Проверил: {supervisor_title}\n{supervisor}' if supervisor_title else f'Проверил:\n{supervisor}'
        run = p.add_run(sup_text)
        run.font.name = gost["font_main"]
        run.font.size = Pt(14)

    # Отступы
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)

    # Город и год
    add_centered(f'{uni["city"]} — {year}', 14)

    doc.add_page_break()

    # === СОДЕРЖАНИЕ ===
    doc.add_heading('СОДЕРЖАНИЕ', level=1)

    toc_items = [
        'Введение',
        '1. Первый раздел',
        '1.1. Подраздел',
        '2. Второй раздел',
        '2.1. Подраздел',
        '3. Третий раздел',
        'Заключение',
        'Список использованных источников',
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.first_line_indent = Cm(0)

    doc.add_page_break()

    # === ВВЕДЕНИЕ ===
    doc.add_heading('ВВЕДЕНИЕ', level=1)
    doc.add_paragraph(
        f'Актуальность темы «{title}» обусловлена тем, что... '
        '[Описать актуальность. Рекомендуемый объём: 1-2 страницы.]'
    )
    p = doc.add_paragraph()
    run = p.add_run('Цель работы')
    run.bold = True
    p.add_run(' — [сформулировать цель реферата].')

    p = doc.add_paragraph()
    run = p.add_run('Задачи работы:')
    run.bold = True

    for task in ['изучить теоретические аспекты...', 'проанализировать...', 'рассмотреть...', 'сделать выводы...']:
        doc.add_paragraph(f'— {task}')

    doc.add_page_break()

    # === ОСНОВНАЯ ЧАСТЬ ===
    doc.add_heading('1. ПЕРВЫЙ РАЗДЕЛ', level=1)
    doc.add_paragraph('[Текст первого раздела. Рекомендуется раскрыть теоретические основы темы.]')
    doc.add_heading('1.1. Подраздел', level=2)
    doc.add_paragraph('[Текст подраздела.]')
    doc.add_page_break()

    doc.add_heading('2. ВТОРОЙ РАЗДЕЛ', level=1)
    doc.add_paragraph('[Текст второго раздела.]')
    doc.add_heading('2.1. Подраздел', level=2)
    doc.add_paragraph('[Текст подраздела.]')
    doc.add_page_break()

    doc.add_heading('3. ТРЕТИЙ РАЗДЕЛ', level=1)
    doc.add_paragraph('[Текст третьего раздела.]')
    doc.add_page_break()

    # === ЗАКЛЮЧЕНИЕ ===
    doc.add_heading('ЗАКЛЮЧЕНИЕ', level=1)
    doc.add_paragraph(
        '[Подвести итоги, сформулировать выводы. Объём: 1-2 страницы.]'
    )
    doc.add_page_break()

    # === СПИСОК ЛИТЕРАТУРЫ ===
    doc.add_heading('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ', level=1)
    refs = [
        'Фамилия, И.О. Название книги / И.О. Фамилия. — М.: Издательство, 2023. — 200 с.',
        'Фамилия, И.О. Название статьи // Журнал. — 2024. — № 1. — С. 10-15.',
        'Название ресурса [Электронный ресурс]. — URL: https://example.com (дата обращения: 01.01.2026).',
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph(f'{i}. {ref}')

    # === СОХРАНЕНИЕ ===
    if not output:
        safe_title = "".join(c if c.isalnum() or c in ' _-' else '' for c in title)[:50].strip()
        output = f"Реферат - {safe_title}.docx"

    doc.save(output)
    print(f"\n  Реферат создан: {output}")
    print(f"  Тема: {title}")
    print(f"  Автор: {author}")
    print(f"  Группа: {group}")
    print(f"  Оформление: ГОСТ 7.32-2017")
    print(f"  Вуз: {UNIVERSITY_INFO['branch']} {UNIVERSITY_INFO['short_name']}")
    return output



# =============================================================================
# CLI ИНТЕРФЕЙС
# =============================================================================

def main():
    parser = argparse.ArgumentParser(
        description='Генератор презентаций и рефератов — Тверской филиал РГУ им. А.Н. Косыгина',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Примеры использования:
  python generator.py presentation "Искусственный интеллект в медицине" --author "Иванов И.И." --group "ИТ-101"
  python generator.py referat "Основы программирования" --author "Петров П.П." --group "ИТ-101" --discipline "Информатика"
  python generator.py presentation "Мода и стиль XX века" --theme art --slides 10
  python generator.py themes
        """
    )
    subparsers = parser.add_subparsers(dest='command', help='Команда')

    # --- Показать темы ---
    subparsers.add_parser('themes', help='Показать доступные темы дизайна')

    # --- Презентация ---
    pres_parser = subparsers.add_parser('presentation', aliases=['pres', 'p'],
                                         help='Создать презентацию (.pptx)')
    pres_parser.add_argument('title', help='Тема презентации')
    pres_parser.add_argument('--author', default='Студент', help='ФИО автора')
    pres_parser.add_argument('--group', default='XX-XXX', help='Номер группы')
    pres_parser.add_argument('--department', default='', help='Кафедра')
    pres_parser.add_argument('--discipline', default='', help='Дисциплина')
    pres_parser.add_argument('--slides', type=int, default=8, help='Количество слайдов (по умолчанию 8)')
    pres_parser.add_argument('--subtitle', default='', help='Подзаголовок')
    pres_parser.add_argument('--theme', default=None,
                              help='Тема дизайна (tech/humanities/science/art/business/minimalist). Если не указана — определяется автоматически.')
    pres_parser.add_argument('--output', '-o', help='Имя выходного файла')

    # --- Реферат ---
    ref_parser = subparsers.add_parser('referat', aliases=['ref', 'r'],
                                        help='Создать реферат (.docx)')
    ref_parser.add_argument('title', help='Тема реферата')
    ref_parser.add_argument('--author', default='Студент', help='ФИО автора')
    ref_parser.add_argument('--group', default='XX-XXX', help='Номер группы')
    ref_parser.add_argument('--department', default='', help='Кафедра')
    ref_parser.add_argument('--discipline', default='', help='Дисциплина')
    ref_parser.add_argument('--supervisor', default='', help='ФИО преподавателя')
    ref_parser.add_argument('--supervisor-title', default='', help='Должность преподавателя')
    ref_parser.add_argument('--output', '-o', help='Имя выходного файла')

    args = parser.parse_args()

    if not args.command:
        parser.print_help()
        print("\n  Подсказка: используйте 'python generator.py themes' чтобы увидеть доступные стили")
        sys.exit(1)

    if args.command == 'themes':
        print("\n  Доступные темы дизайна презентаций:")
        print("  " + "=" * 50)
        for key, name in list_themes():
            print(f"    {key:<12} — {name}")
        print("\n  Если тема не указана, она определяется автоматически по названию.")
        print("  Пример: python generator.py p \"Нейросети\" --theme tech\n")
        sys.exit(0)

    if args.command in ('presentation', 'pres', 'p'):
        create_presentation(
            title=args.title,
            author=args.author,
            group=args.group,
            department=args.department,
            discipline=args.discipline,
            slides_count=args.slides,
            subtitle=args.subtitle,
            theme_key=args.theme,
            output=args.output,
        )
    elif args.command in ('referat', 'ref', 'r'):
        create_referat(
            title=args.title,
            author=args.author,
            group=args.group,
            department=args.department,
            discipline=args.discipline,
            supervisor=args.supervisor,
            supervisor_title=getattr(args, 'supervisor_title', ''),
            output=args.output,
        )


if __name__ == "__main__":
    main()
